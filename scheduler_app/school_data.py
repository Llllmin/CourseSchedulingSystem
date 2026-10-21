"""Validated data editing and atomic Excel import, using stable school codes."""

import csv
import secrets
import sqlite3
from collections import Counter
from contextlib import closing
from datetime import datetime
from pathlib import Path

from scheduler_app.auth import hash_password
from scheduler_app.config import DAYS, PERIODS, WEEKLY_PERIODS
from scheduler_app.database import connect

# Workbook headers also drive the administrator's editable forms.
FIELDS = {
    "Students": ("student_code", "name", "grade"),
    "Teachers": ("teacher_code", "name", "department"),
    "Rooms": ("room_code", "capacity", "room_type"),
    "Courses": ("course_code", "name", "requires_room_type"),
    "Sections": ("section_code", "course_code", "teacher_code", "grade", "level", "duration"),
    "Enrollments": ("student_code", "section_code"),
    "TeacherAvailability": ("teacher_code", "day", "period", "available"),
    "BlockedPeriods": ("day", "period", "reason"),
}
KEYS = {name: fields[:1] for name, fields in FIELDS.items()}
KEYS.update(Enrollments=("student_code", "section_code"),
            TeacherAvailability=("teacher_code", "day", "period"),
            BlockedPeriods=("day", "period"))


def snapshot(db_path):
    """Export relational IDs as readable codes, not database internals."""
    queries = {
        "Students": "SELECT student_code, name, grade FROM students ORDER BY student_code",
        "Teachers": "SELECT teacher_code, name, department FROM teachers ORDER BY teacher_code",
        "Rooms": "SELECT room_code, capacity, room_type FROM rooms ORDER BY room_code",
        "Courses": "SELECT course_code, name, requires_room_type FROM courses ORDER BY course_code",
        "Sections": """SELECT s.section_code, c.course_code, t.teacher_code, s.grade, s.level, s.duration
            FROM course_sections s JOIN courses c ON c.id=s.course_id JOIN teachers t ON t.id=s.teacher_id
            ORDER BY s.section_code""",
        "Enrollments": """SELECT st.student_code, s.section_code FROM student_enrollments e
            JOIN students st ON st.id=e.student_id JOIN course_sections s ON s.id=e.section_id
            ORDER BY st.student_code,s.section_code""",
        "TeacherAvailability": """SELECT t.teacher_code,a.day,a.period,a.available FROM teacher_availability a
            JOIN teachers t ON t.id=a.teacher_id ORDER BY t.teacher_code,a.day,a.period""",
        "BlockedPeriods": "SELECT day,period,reason FROM blocked_periods ORDER BY day,period",
    }
    with connect(db_path) as conn:
        return {name: [dict(r) for r in conn.execute(query)] for name, query in queries.items()}


def validate_data(data, strict=True):
    """Return a normalized copy or fail with a sheet and row-specific message."""
    clean = {}
    indexes = {}
    for sheet, fields in FIELDS.items():
        if sheet not in data:
            raise ValueError(f"Missing sheet: {sheet}")
        clean[sheet], indexes[sheet] = [], {}
        for number, original in enumerate(data[sheet], 2):
            row = {field: str(original.get(field, "")).strip() if original.get(field) is not None else "" for field in fields}
            prefix = f"{sheet} row {number}: "
            if any(not value for value in row.values()):
                raise ValueError(prefix + "all columns are required")
            for field in ("period", "duration", "capacity", "available"):
                if field in row:
                    try:
                        value = float(row[field])
                        if not value.is_integer():
                            raise ValueError()
                        row[field] = int(value)
                    except ValueError:
                        raise ValueError(prefix + f"{field} must be an integer") from None
            if "day" in row and row["day"] not in DAYS:
                raise ValueError(prefix + "day must be Monday-Friday")
            if "period" in row and row["period"] not in PERIODS:
                raise ValueError(prefix + "period must be 1-11")
            if "available" in row and row["available"] not in (0, 1):
                raise ValueError(prefix + "available must be 0 or 1")
            if "capacity" in row and row["capacity"] < 1:
                raise ValueError(prefix + "capacity must be positive")
            for field in ("room_type", "requires_room_type"):
                if field in row and row[field] not in ("standard", "lab"):
                    raise ValueError(prefix + f"{field} must be standard or lab")
            if sheet == "Sections":
                if row["level"] not in WEEKLY_PERIODS:
                    raise ValueError(prefix + "level must be HL, SL, PE, TOK or EE")
                if row["duration"] not in (1, 2) or WEEKLY_PERIODS[row["level"]] % row["duration"]:
                    raise ValueError(prefix + "duration must be 1 or 2 and divide the weekly period count")
            key = tuple(row[field] for field in KEYS[sheet])
            if key in indexes[sheet]:
                raise ValueError(prefix + f"duplicate key {key}")
            indexes[sheet][key] = row
            clean[sheet].append(row)

    # Resolve references only after all worksheets have been normalized.
    references = {"student_code": "Students", "teacher_code": "Teachers",
                  "course_code": "Courses", "section_code": "Sections"}
    for sheet in ("Sections", "Enrollments", "TeacherAvailability"):
        for number, row in enumerate(clean[sheet], 2):
            for field, target in references.items():
                if field in row and field not in KEYS[sheet] and (row[field],) not in indexes[target]:
                    raise ValueError(f"{sheet} row {number}: unknown {field} {row[field]}")
                # Enrollment fields are both key columns but still foreign keys.
                if sheet == "Enrollments" and field in row and (row[field],) not in indexes[target]:
                    raise ValueError(f"{sheet} row {number}: unknown {field} {row[field]}")
            if sheet == "TeacherAvailability" and (row["teacher_code"],) not in indexes["Teachers"]:
                raise ValueError(f"{sheet} row {number}: unknown teacher {row['teacher_code']}")
    enrollments = {r["student_code"]: [] for r in clean["Students"]}
    section_sizes = Counter()
    for row in clean["Enrollments"]:
        section = indexes["Sections"][(row["section_code"],)]
        student = indexes["Students"][(row["student_code"],)]
        if student["grade"] != section["grade"]:
            raise ValueError(f"Enrollments: {row['student_code']} grade does not match {row['section_code']}")
        enrollments[row["student_code"]].append(section)
        section_sizes[row["section_code"]] += 1
    for code, sections in enrollments.items():
        courses = [s["course_code"] for s in sections]
        if len(set(courses)) != len(courses):
            raise ValueError(f"Student {code}: duplicate subject enrollment")
        counts = Counter(s["level"] for s in sections)
        if strict and counts != Counter(HL=3, SL=3, PE=1, TOK=1, EE=1):
            raise ValueError(f"Student {code}: requires 3 HL + 3 SL + one PE, TOK and EE section; got {dict(counts)}")
    if strict:
        if not clean["Students"] or not clean["Sections"]:
            raise ValueError("Students and Sections cannot be empty")
        for section in clean["Sections"]:
            course = indexes["Courses"][(section["course_code"],)]
            if not any(r["capacity"] >= section_sizes[section["section_code"]] and
                       (course["requires_room_type"] == "standard" or r["room_type"] == "lab")
                       for r in clean["Rooms"]):
                raise ValueError(f"Section {section['section_code']}: no room has the required type and capacity")
    return clean


def backup_database(db_path):
    """Use SQLite's backup API so the snapshot is consistent."""
    folder = Path(db_path).resolve().parent / "backups"
    folder.mkdir(exist_ok=True)
    target = folder / f"school_{datetime.now():%Y%m%d_%H%M%S_%f}.db"
    with closing(sqlite3.connect(db_path)) as source, closing(sqlite3.connect(target)) as destination:
        source.backup(destination)
    return target


def replace_data(db_path, data, strict=True):
    """Validate first, back up, then replace all school records in one transaction.

    Existing accounts follow stable codes. New people receive random passwords;
    credentials are returned to the caller, never included in a source export.
    """
    data = validate_data(data, strict)
    backup = backup_database(db_path)
    credentials = []
    with connect(db_path) as conn:
        accounts = {}
        for role, table, field in (("student", "students", "student_code"), ("teacher", "teachers", "teacher_code")):
            for row in conn.execute(f"SELECT u.*,p.{field} AS code FROM users u JOIN {table} p ON p.id=u.{role}_id WHERE u.role=?", (role,)):
                accounts[(role, row["code"])] = dict(row)
        for table in ("schedule_assignments", "student_enrollments", "teacher_availability", "course_sections", "courses", "rooms", "students", "teachers", "blocked_periods"):
            conn.execute(f"DELETE FROM {table}")
        ids = {}
        for sheet, table in (("Students", "students"), ("Teachers", "teachers"), ("Rooms", "rooms"), ("Courses", "courses")):
            ids[sheet] = {}
            fields = FIELDS[sheet]
            for row in data[sheet]:
                cur = conn.execute(f"INSERT INTO {table}({','.join(fields)}) VALUES ({','.join('?' for _ in fields)})", tuple(row[f] for f in fields))
                ids[sheet][row[fields[0]]] = cur.lastrowid
        ids["Sections"] = {}
        for r in data["Sections"]:
            cur = conn.execute("INSERT INTO course_sections(section_code,course_id,teacher_id,grade,level,duration) VALUES (?,?,?,?,?,?)",
                               (r["section_code"], ids["Courses"][r["course_code"]], ids["Teachers"][r["teacher_code"]], r["grade"], r["level"], r["duration"]))
            ids["Sections"][r["section_code"]] = cur.lastrowid
        for r in data["Enrollments"]:
            conn.execute("INSERT INTO student_enrollments VALUES (?,?)", (ids["Students"][r["student_code"]], ids["Sections"][r["section_code"]]))
        for r in data["TeacherAvailability"]:
            conn.execute("INSERT INTO teacher_availability VALUES (?,?,?,?)", (ids["Teachers"][r["teacher_code"]], r["day"], r["period"], r["available"]))
        for r in data["BlockedPeriods"]:
            conn.execute("INSERT INTO blocked_periods VALUES (?,?,?)", (r["day"], r["period"], r["reason"]))
        for role, sheet in (("student", "Students"), ("teacher", "Teachers")):
            for code, person_id in ids[sheet].items():
                old = accounts.get((role, code))
                username = old["username"] if old else f"{role}_{code}"
                password = secrets.token_urlsafe(12)
                password_hash = old["password_hash"] if old else hash_password(password)
                conn.execute(f"INSERT INTO users(username,password_hash,role,{role}_id) VALUES (?,?,?,?)", (username, password_hash, role, person_id))
                if not old:
                    credentials.append((role, code, username, password))
    return backup, credentials


def save_credentials(db_path, credentials):
    if not credentials:
        return None
    folder = Path(db_path).resolve().parent / "data"
    folder.mkdir(exist_ok=True)
    target = folder / f"new_accounts_{datetime.now():%Y%m%d_%H%M%S_%f}.csv"
    with target.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(("role", "code", "username", "password"))
        writer.writerows(credentials)
    target.chmod(0o600)
    return target


def export_workbook(path, data=None):
    """Create a blank template or export current school data (no passwords)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    for sheet, fields in FIELDS.items():
        page = book.create_sheet(sheet)
        page.append(fields)
        for row in (data or {}).get(sheet, []):
            page.append([row[field] for field in fields])
        page.freeze_panes = "A2"
        page.auto_filter.ref = page.dimensions
        for cell in page[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="285E7C")
            page.column_dimensions[cell.column_letter].width = max(22, len(str(cell.value)) + 4)
    book.save(path)
    book.close()


def read_workbook(path):
    """Reject formulas and mismatched headers rather than silently losing data."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        data = {}
        for sheet, fields in FIELDS.items():
            if sheet not in book.sheetnames:
                raise ValueError(f"Missing sheet: {sheet}")
            rows = book[sheet].iter_rows(values_only=True)
            header = next(rows, ())
            if tuple(header) != fields:
                raise ValueError(f"{sheet}: expected columns {', '.join(fields)}")
            data[sheet] = []
            for number, values in enumerate(rows, 2):
                if all(v is None for v in values):
                    continue
                if any(isinstance(v, str) and v.startswith('=') for v in values):
                    raise ValueError(f"{sheet} row {number}: formulas are not allowed")
                data[sheet].append(dict(zip(fields, values)))
        return validate_data(data)
    finally:
        book.close()
