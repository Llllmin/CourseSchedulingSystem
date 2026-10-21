"""Integration coverage for weekly loads, atomic imports and persistent edits."""

import copy
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from scheduler_app.config import DAYS, WEEKLY_PERIODS
from scheduler_app.database import connect, initialize_database
from scheduler_app.scheduler import TimetableScheduler, timetable_rows
from scheduler_app.school_data import (export_workbook, read_workbook, replace_data,
                                       snapshot, validate_data)
from scheduler_app.seed_data import demo_data, seed_demo_data


class SchoolDataTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.path = str(Path(self.folder.name) / "school.db")
        initialize_database(self.path)
        seed_demo_data(self.path)
        self.scheduler = TimetableScheduler(self.path)

    def tearDown(self):
        self.folder.cleanup()

    def test_full_week_and_all_conflicts(self):
        self.assertTrue(self.scheduler.generate()[0])
        rows = timetable_rows(self.path)
        self.assertEqual(len(rows), 102)
        counts = Counter(r["section_id"] for r in rows)
        for row in rows:
            self.assertEqual(counts[row["section_id"]] * row["duration"], WEEKLY_PERIODS[row["level"]])
            self.assertTrue(self.scheduler.validate_assignment(row["section_id"], row["day"], row["period"], row["room_id"], row["occurrence"]).valid)
        for student in range(1, 61):
            lessons = timetable_rows(self.path, "student", student_id=student)
            self.assertEqual(sum(r["duration"] for r in lessons), 34)
            self.assertEqual(set(r["day"] for r in lessons), set(DAYS))
        with connect(self.path) as conn:
            section_teacher = {r["id"]: r["teacher_id"] for r in conn.execute("SELECT * FROM course_sections")}
        for teacher in range(1, 21):
            lessons = timetable_rows(self.path, "teacher", teacher_id=teacher)
            self.assertTrue(all(section_teacher[r["section_id"]] == teacher for r in lessons))

    def test_excel_roundtrip_and_persistence(self):
        path = Path(self.folder.name) / "school.xlsx"
        before = snapshot(self.path)
        export_workbook(path, before)
        data = read_workbook(path)
        data["Students"][0]["name"] = "Updated student"
        backup, accounts = replace_data(self.path, data)
        self.assertTrue(backup.exists())
        self.assertEqual(accounts, [])
        initialize_database(self.path)
        self.assertEqual(snapshot(self.path)["Students"][0]["name"], "Updated student")
        with connect(self.path) as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM users WHERE username='student001'").fetchone()[0], 1)
        with connect(str(backup)) as conn:
            self.assertEqual(conn.execute("SELECT name FROM students WHERE student_code='S001'").fetchone()[0], "Student 001")

    def test_invalid_import_leaves_existing_data_and_schedule(self):
        self.assertTrue(self.scheduler.generate()[0])
        before = snapshot(self.path)
        placements = timetable_rows(self.path)
        invalid = copy.deepcopy(before)
        invalid["Enrollments"].pop(0)
        with self.assertRaisesRegex(ValueError, "3 HL"):
            replace_data(self.path, invalid)
        self.assertEqual(snapshot(self.path), before)
        self.assertEqual(timetable_rows(self.path), placements)

    def test_validation_duplicate_missing_reference_capacity_and_grade(self):
        for mutate, message in (
            (lambda d: d["Courses"].append(d["Courses"][0]), "duplicate key"),
            (lambda d: d["Sections"][0].update(teacher_code="missing"), "unknown teacher_code"),
            (lambda d: [r.update(capacity=1) for r in d["Rooms"]], "no room"),
            (lambda d: d["Sections"][0].update(grade="G99"), "grade does not match"),
        ):
            data = demo_data()
            mutate(data)
            with self.assertRaisesRegex(ValueError, message):
                validate_data(data)

    def test_transaction_rolls_back_on_username_collision(self):
        before = snapshot(self.path)
        data = copy.deepcopy(before)
        # A new student's generated username would collide with an admin account.
        with connect(self.path) as conn:
            conn.execute("INSERT INTO users(username,password_hash,role) VALUES ('student_NEW','x','admin')")
        data["Students"].append(dict(student_code="NEW", name="New", grade="G12"))
        with self.assertRaises(Exception):
            replace_data(self.path, data, strict=False)
        self.assertEqual(snapshot(self.path), before)

    def test_new_accounts_and_edits_clear_stale_schedule(self):
        self.assertTrue(self.scheduler.generate()[0])
        data = snapshot(self.path)
        data["Students"].append(dict(student_code="NEW", name="New", grade="G12"))
        _, credentials = replace_data(self.path, data, strict=False)
        self.assertEqual(len(credentials), 1)
        self.assertEqual(credentials[0][2], "student_NEW")
        self.assertGreaterEqual(len(credentials[0][3]), 12)
        self.assertEqual(timetable_rows(self.path), [])
        # Incomplete drafts are editable, but cannot be scheduled or imported.
        self.assertFalse(self.scheduler.generate()[0])

    def test_failed_generation_preserves_previous_schedule(self):
        self.assertTrue(self.scheduler.generate()[0])
        before = timetable_rows(self.path)
        with connect(self.path) as conn:
            conn.executemany("INSERT INTO blocked_periods VALUES (?,?,?)",
                             [(day, period, "Closed") for day in DAYS for period in range(1, 12)])
        ok, message = self.scheduler.generate()
        self.assertFalse(ok)
        self.assertIn("No feasible slots", message)
        self.assertEqual(timetable_rows(self.path), before)

    def test_second_occurrence_edit_does_not_replace_first(self):
        self.assertTrue(self.scheduler.save_manual_assignment(1, "Monday", 1, 1, 1).valid)
        self.assertTrue(self.scheduler.save_manual_assignment(1, "Tuesday", 1, 1, 2).valid)
        self.assertFalse(self.scheduler.validate_assignment(1, "Monday", 1, 1, 2).valid)
        self.assertTrue(self.scheduler.save_manual_assignment(1, "Wednesday", 2, 1, 2).valid)
        rows = timetable_rows(self.path)
        self.assertEqual(len(rows), 2)
        self.assertEqual(next(r["day"] for r in rows if r["occurrence"] == 1), "Monday")

    def test_header_and_formula_rejected(self):
        from openpyxl import load_workbook
        path = Path(self.folder.name) / "bad.xlsx"
        export_workbook(path, demo_data())
        book = load_workbook(path)
        book["Students"]["A1"] = "incorrect"
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, "expected columns"):
            read_workbook(path)
        export_workbook(path, demo_data())
        book = load_workbook(path)
        book["Students"]["B2"] = '=1+1'
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, "formulas"):
            read_workbook(path)

    def test_double_lessons_have_exact_weekly_counts(self):
        data = snapshot(self.path)
        for section in data["Sections"]:
            if section["level"] in ("HL", "SL"):
                section["duration"] = 2
        replace_data(self.path, data)
        ok, message = self.scheduler.generate()
        self.assertTrue(ok, message)
        for student in range(1, 61):
            self.assertEqual(sum(r["duration"] for r in timetable_rows(self.path, "student", student_id=student)), 34)

    def test_old_assignment_schema_migrates_without_loss(self):
        with connect(self.path) as conn:
            conn.execute("DROP TABLE schedule_assignments")
            conn.execute("""CREATE TABLE schedule_assignments (
                id INTEGER PRIMARY KEY, section_id INTEGER NOT NULL UNIQUE,
                day TEXT NOT NULL, period INTEGER NOT NULL, room_id INTEGER NOT NULL,
                FOREIGN KEY(section_id) REFERENCES course_sections(id) ON DELETE CASCADE,
                FOREIGN KEY(room_id) REFERENCES rooms(id) ON DELETE RESTRICT)""")
            conn.execute("INSERT INTO schedule_assignments VALUES (99,1,'Monday',1,1)")
        initialize_database(self.path)
        initialize_database(self.path)
        with connect(self.path) as conn:
            row = conn.execute("SELECT * FROM schedule_assignments WHERE id=99").fetchone()
            self.assertEqual(row["occurrence"], 1)
            self.assertEqual(row["section_id"], 1)
            conn.execute("INSERT INTO schedule_assignments(section_id,occurrence,day,period,room_id) VALUES (1,2,'Tuesday',1,1)")

    def test_unavailable_teacher_and_daily_spread_limit(self):
        unavailable = self.scheduler.validate_assignment(1, "Wednesday", 1, 1)
        self.assertFalse(unavailable.valid)
        self.assertIn("Teacher", unavailable.message)
        self.assertTrue(self.scheduler.save_manual_assignment(1, "Monday", 1, 1, 1).valid)
        self.assertTrue(self.scheduler.save_manual_assignment(1, "Monday", 2, 1, 2).valid)
        result = self.scheduler.validate_assignment(1, "Monday", 3, 1, 3)
        self.assertFalse(result.valid)
        self.assertIn("at most 2", result.message)


if __name__ == "__main__":
    unittest.main()
